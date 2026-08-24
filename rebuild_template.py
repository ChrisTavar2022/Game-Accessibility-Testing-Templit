"""Rebuild the Game Accessibility Testing template.

Idempotent: safe to re-run against its own output.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

SRC = r"Game Accessibility Testing templit.xlsx"
AUTHOR_NAME = "Christopher Tavarez"

wb = openpyxl.load_workbook(SRC)
# Names from superseded designs; harmless if absent.
for stale in ("SummaryTabNames", "SeverityNumbers"):
    wb.defined_names.pop(stale, None)

HEADER_FILL = PatternFill("solid", fgColor="FF38761D")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
BODY_ALIGN = Alignment(horizontal="left", vertical="top", wrap_text=True)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", size=10, bold=True)
LOCKED = Protection(locked=True)
UNLOCKED = Protection(locked=False)
COMMENT_AUTHOR = "Accessibility Template"

DATA_ROWS = 2000
FIRST_DATA_ROW = 2
LAST_DATA_ROW = FIRST_DATA_ROW + DATA_ROWS - 1  # 2001

# The value itself carries the label so a screen reader announces "1 - Cosmetic"
# rather than a bare digit whose meaning sits in a column the dropdown never reads.
SEVERITY_SEED = [
    ("1 - Cosmetic", "Visual or audio polish only; does not affect play or comprehension."),
    ("2 - Minor", "Noticeable, but there is an easy workaround and play continues normally."),
    ("3 - Moderate", "A real barrier; a workaround exists but is awkward, slow, or easy to miss."),
    ("4 - High", "Major barrier; some content or features are effectively unusable."),
    ("5 - Blocking", "Prevents progress entirely; the player cannot continue."),
]
SEVERITY_ROW_COUNT = 15  # 5 seeded + room for longer scales (e.g. WCAG / GAG wording)


def protect_sheet(ws):
    """Lock headers/formulas while leaving UNLOCKED cells editable."""
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False
    ws.protection.insertRows = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def reset_sheet(name, legacy_names=()):
    """Return a genuinely blank sheet, dropping any existing version outright.

    Clearing cells in place leaves orphaned styling behind (fills, fonts, locked
    protection) from earlier layouts, which inflates the used range and strands
    coloured header cells with no text. Deleting and recreating guarantees the
    rebuild is clean and repeatable.
    """
    for existing in (name,) + tuple(legacy_names):
        if existing in wb.sheetnames:
            wb.remove(wb[existing])
    return wb.create_sheet(name)


def write_header(ws, row, labels, start_col=1):
    for offset, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + offset, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = HEADER_ALIGN
        c.protection = LOCKED


def add_table(ws, name, ref):
    tbl = Table(displayName=name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight1", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tbl)


def write_instructions(ws, lines):
    """Write a title + numbered guidance block. Returns the next free row."""
    for i, text in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = TITLE_FONT if i == 1 else BODY_FONT
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c.protection = LOCKED
        ws.row_dimensions[i].height = 20 if i == 1 else 30
    return len(lines) + 2  # leave one blank separator row


# ===========================================================================
# 1. Setup sheet
# ===========================================================================
setup = reset_sheet("Setup", legacy_names=("Accessibilitty testing log",))

row_cursor = write_instructions(setup, [
    "How to Use This Workbook",
    "1. Fill in your game/mod info in the table further down this sheet.",
    "2. Go to the 'Game Sections' tab and list every section/component of your game (Title Screen, Gameplay, Pause Menu, etc.) - the Game Section dropdown on every log tab updates automatically as you add rows there. Add 3 or 30 - there's no limit.",
    "3. Log one issue per row in the Visual, Screen Reader, Keyboard, and Controller tabs. Free-type the description fields; use the dropdowns for Game Section, Error Type, Severity, Platform, Input Method, Screen Reader Used, and Status. The Error Type lists below are yours to edit - for example, swap in terminology from a framework like WCAG or the Game Accessibility Guidelines, since those categorize the kind of issue rather than how severe it is.",
    "4. NVDA/JAWS tip: to choose from a dropdown cell, select the cell but do NOT press F2. Press Alt+Down Arrow to open its list, use Up/Down Arrow to move through the options, then Enter to select. F2 switches to text-editing mode, where Up/Down Arrow won't respond to the list - if that happens, press Escape first, then use Alt+Down Arrow instead.",
    "5. Status is always a plain text word (Open, In Progress, Resolved, Won't Fix, Duplicate, Needs More Info) - never rely on color alone.",
    "6. Severity levels live on the 'Severity Scale' tab and are fully editable - rename them, add more, or replace them with your own wording. The Severity dropdown and the Summary tab follow whatever you put there.",
    "7. Header rows on most tabs are locked to prevent accidental edits. All the entry cells below them stay fully editable. The 'Severity Scale' tab is left completely unlocked so you can restructure it freely.",
    "8. Each row's Issue ID appears automatically as soon as you pick a Game Section for that row. Rows with no Game Section chosen stay completely blank.",
    "9. Each log tab has 2,000 ready-to-use rows, so you won't run out for any realistic number of bugs. If you ever do fill all 2,000: go to Review > Unprotect Sheet (there is no password), then just start typing in the first empty row underneath - the table and its dropdowns extend automatically. The Issue ID does not copy itself, so also copy the Issue ID cell from the row above into the new row. Re-protect the sheet afterwards if you want the headers locked again.",
    "10. See the 'Summary' tab for issue totals by severity and by tab, plus instructions for adding your own custom log tabs.",
    f"Original template created by {AUTHOR_NAME}. Please keep this credit if you share or adapt it.",
])

# --- Game / Mod Information ---
setup.cell(row=row_cursor, column=1, value="Game / Mod Information").font = TITLE_FONT
setup.cell(row=row_cursor, column=1).protection = LOCKED
row_cursor += 1

info_header_row = row_cursor
write_header(setup, info_header_row, ["Field", "Value"])
info_fields = ["Name of Game", "Author / Mod Name", "Current Mod Version", "Date Created"]
for i, label in enumerate(info_fields):
    r = info_header_row + 1 + i
    lc = setup.cell(row=r, column=1, value=label)
    lc.font = BODY_FONT
    lc.alignment = Alignment(vertical="top")
    lc.protection = LOCKED
    vc = setup.cell(row=r, column=2)
    vc.font = BODY_FONT
    vc.protection = UNLOCKED

info_table_last_row = info_header_row + len(info_fields)
add_table(setup, "tbl_GameInfo", f"A{info_header_row}:B{info_table_last_row}")
row_cursor = info_table_last_row + 2

# --- Dropdown source lists ---
setup.cell(row=row_cursor, column=1,
           value="Dropdown Lists (add or edit rows below each heading; dropdowns update automatically)").font = TITLE_FONT
setup.cell(row=row_cursor, column=1).protection = LOCKED

LIST_HEADER_ROW = row_cursor + 1
# (title, seed values, rows reserved below the header)
# Game Sections and Severity levels live on their own sheets - both need far more
# room to grow. WCAG / Game Accessibility Guidelines terminology belongs on the
# Error Type lists below, not Severity - those categorize the kind of issue, not
# how severe it is.
lists = {
    "A": ("Status", [
        "Open", "In Progress", "Resolved", "Won't Fix", "Duplicate", "Needs More Info",
    ], 9),
    "B": ("Platform", [
        "PC (Windows)", "PC (Steam Deck)", "PlayStation 5", "PlayStation 4",
        "Xbox Series X|S", "Xbox One", "Nintendo Switch", "Nintendo Switch 2",
        "Mac", "Linux", "Android", "Other",
    ], 13),
    "C": ("Input Method", [
        "Keyboard", "Mouse", "Keyboard + Mouse", "Xbox Controller",
        "PlayStation Controller", "Switch Pro Controller", "Touch",
        "Switch / Adaptive Controller", "Other",
    ], 11),
    "D": ("Screen Reader Used", [
        "NVDA", "JAWS", "Narrator", "VoiceOver", "TalkBack", "None / N/A", "Other",
    ], 9),
    "E": ("Visual Error Types", [
        "Color-Only Indicator", "Text Too Small / Not Scalable", "UI Element Not Visible",
        "Font / Rendering Issue", "Flashing / Strobing Content", "Missing Visual Feedback", "Other",
    ], 15),
    "F": ("Screen Reader Error Types", [
        "Element Not Announced", "Incorrect / Missing Label", "Focus Not Set / Lost",
        "Reading Order Incorrect", "Silent on Interaction", "Redundant / Verbose Announcement",
        "Live Region Not Announced", "Other",
    ], 15),
    "G": ("Keyboard Error Types", [
        "Not Keyboard Accessible", "Keyboard Trap", "Focus Indicator Missing / Unclear",
        "Tab Order Incorrect", "No Remap / Rebind Option", "Shortcut Conflict", "Other",
    ], 15),
    "H": ("Controller Error Types", [
        "Button Not Mapped", "No Remap / Rebind Option", "Prompt / Glyph Mismatch",
        "Vibration / Haptic Issue", "Input Not Recognized", "Other",
    ], 15),
}

named_range_targets = {}
for col_letter, (title, values, row_count) in lists.items():
    last_row = LIST_HEADER_ROW + row_count
    hc = setup[f"{col_letter}{LIST_HEADER_ROW}"]
    hc.value = title
    hc.fill = HEADER_FILL
    hc.font = HEADER_FONT
    hc.alignment = HEADER_ALIGN
    hc.protection = LOCKED
    for i in range(row_count):
        cell = setup[f"{col_letter}{LIST_HEADER_ROW + 1 + i}"]
        cell.value = values[i] if i < len(values) else None
        cell.font = BODY_FONT
        cell.protection = UNLOCKED
    named_range_targets[title] = f"${col_letter}${LIST_HEADER_ROW + 1}:${col_letter}${last_row}"

# Captured for the Summary tab's Issues by Status block, which mirrors the Status
# list the same way Issues by Severity mirrors the 'Severity Scale' sheet.
STATUS_FIRST_DATA_ROW = LIST_HEADER_ROW + 1
STATUS_ROW_COUNT = lists["A"][2]

for col, w in {"A": 26, "B": 20, "C": 24, "D": 20, "E": 26, "F": 28, "G": 26, "H": 24}.items():
    setup.column_dimensions[col].width = w
setup.freeze_panes = "A2"
protect_sheet(setup)

for defined_name, title in {
    "StatusList": "Status",
    "PlatformList": "Platform",
    "InputMethodList": "Input Method",
    "ScreenReaderList": "Screen Reader Used",
    "VisualErrorTypes": "Visual Error Types",
    "ScreenReaderErrorTypes": "Screen Reader Error Types",
    "KeyboardErrorTypes": "Keyboard Error Types",
    "ControllerErrorTypes": "Controller Error Types",
}.items():
    wb.defined_names[defined_name] = DefinedName(
        defined_name, attr_text=f"'Setup'!{named_range_targets[title]}"
    )

# ===========================================================================
# 2. Game Sections sheet - per-game component list feeding the Game Section dropdown
# ===========================================================================
gs = reset_sheet("Game Sections")
gs_instr = gs.cell(row=1, column=1, value=(
    "List every section/component of this game or mod being tested - one per row. "
    "The Game Section dropdown on every log tab updates automatically as you add rows. "
    "500 rows are ready to use; if you ever fill all of them, right-click a row number "
    "in this list and choose Insert to add more - it inherits the same formatting."
))
gs_instr.font = TITLE_FONT
gs_instr.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
gs_instr.protection = LOCKED
gs.row_dimensions[1].height = 30

GS_HEADER_ROW = 2
write_header(gs, GS_HEADER_ROW, ["Game Section"])
GS_SEED = [
    "Title Screen", "Loading Screen", "Main Menu", "Settings Menu",
    "Gameplay", "Pause Menu", "Inventory / Menus", "Dialogue / Cutscenes",
    "Credits", "Other",
]
GS_ROW_COUNT = 500
GS_FIRST_DATA_ROW = GS_HEADER_ROW + 1
GS_LAST_DATA_ROW = GS_HEADER_ROW + GS_ROW_COUNT
for i in range(GS_ROW_COUNT):
    cell = gs.cell(row=GS_FIRST_DATA_ROW + i, column=1,
                   value=GS_SEED[i] if i < len(GS_SEED) else None)
    cell.font = BODY_FONT
    cell.protection = UNLOCKED

gs.column_dimensions["A"].width = 34
add_table(gs, "tbl_GameSections", f"A{GS_HEADER_ROW}:A{GS_LAST_DATA_ROW}")
gs.freeze_panes = f"A{GS_FIRST_DATA_ROW}"
protect_sheet(gs)

wb.defined_names["GameSections"] = DefinedName(
    "GameSections",
    attr_text=f"'Game Sections'!$A${GS_FIRST_DATA_ROW}:$A${GS_LAST_DATA_ROW}",
)

# ===========================================================================
# 3. Severity Scale sheet - deliberately left fully unlocked so teams can rename
#    or restructure the severity levels themselves. WCAG / Game Accessibility
#    Guidelines terminology belongs on the Setup tab's Error Type dropdowns
#    instead - those categorize the kind of issue, not its severity.
# ===========================================================================
sev = reset_sheet("Severity Scale")
# Row numbers are derived, not hardcoded, so the on-sheet guidance can't drift
# out of sync with the layout when the instruction block changes length.
_sev_instr_count = 4
_sev_first_data_row = _sev_instr_count + 2  # instructions, then header, then data
sev_instr_lines = [
    "Issue Severity Scale",
    "This tab is fully unlocked and yours to change. Rename levels, add or remove them, or replace the whole scale with your own wording to match how your team talks about severity. (Looking to use terminology from a framework like WCAG or the Game Accessibility Guidelines instead? Those categorize the kind of issue, not how severe it is, so they fit better on the Error Type dropdown lists on the Setup tab.)",
    "Whatever you put in the 'Severity Value' column becomes the Severity dropdown on every log tab, and the Summary tab totals follow it automatically. Values can be numbers or text.",
    f"Keep entries in a single unbroken run starting at row {_sev_first_data_row} - a blank row in the middle will cut the dropdown short. Use the 'Meaning / Description' column for your own definition of each level.",
]
assert len(sev_instr_lines) == _sev_instr_count, "update _sev_instr_count to match"
for i, text in enumerate(sev_instr_lines, start=1):
    c = sev.cell(row=i, column=1, value=text)
    c.font = TITLE_FONT if i == 1 else BODY_FONT
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c.protection = UNLOCKED
    sev.row_dimensions[i].height = 20 if i == 1 else 34

SEV_HEADER_ROW = len(sev_instr_lines) + 1
write_header(sev, SEV_HEADER_ROW, ["Severity Value", "Meaning / Description"])
for c in (sev.cell(row=SEV_HEADER_ROW, column=1), sev.cell(row=SEV_HEADER_ROW, column=2)):
    c.protection = UNLOCKED  # sheet is unprotected anyway; keep it consistent

SEV_FIRST_DATA_ROW = SEV_HEADER_ROW + 1
SEV_LAST_DATA_ROW = SEV_HEADER_ROW + SEVERITY_ROW_COUNT
for i in range(SEVERITY_ROW_COUNT):
    r = SEV_FIRST_DATA_ROW + i
    value, meaning = SEVERITY_SEED[i] if i < len(SEVERITY_SEED) else (None, None)
    vc = sev.cell(row=r, column=1, value=value)
    vc.font = BODY_FONT
    vc.protection = UNLOCKED
    mc = sev.cell(row=r, column=2, value=meaning)
    mc.font = BODY_FONT
    mc.alignment = BODY_ALIGN
    mc.protection = UNLOCKED

sev.column_dimensions["A"].width = 18
sev.column_dimensions["B"].width = 52
add_table(sev, "tbl_SeverityScale", f"A{SEV_HEADER_ROW}:B{SEV_LAST_DATA_ROW}")
sev.freeze_panes = f"A{SEV_FIRST_DATA_ROW}"
# NOTE: no protect_sheet() here - this tab stays completely unlocked by design.

wb.defined_names["SeverityLevels"] = DefinedName(
    "SeverityLevels",
    attr_text=f"'Severity Scale'!$A${SEV_FIRST_DATA_ROW}:$A${SEV_LAST_DATA_ROW}",
)

# ===========================================================================
# 4. Log sheets
# ===========================================================================
COLUMNS = [
    ("Issue ID", 10, "Auto-generated once you pick a Game Section for this row. Locked - do not edit. Stays blank on unused rows."),
    ("Date Reported", 14, "New field: date this issue was logged."),
    ("Game Section", 18, "Original header: \"Which specific part of the game are you testing?\" Add new options on the 'Game Sections' tab."),
    ("Error Type", 24, "Original header: \"What type of error did you find?\""),
    ("Actual Behavior", 30, "Original header: \"What specific error did you find?\" - describe what actually happens."),
    ("Expected Behavior", 30, "New field: describe what should happen instead (e.g. what a sighted/mouse user experiences, or what the screen reader should announce)."),
    ("Steps to Reproduce", 34, "Original header: \"Describe how to duplicate the problem.\""),
    ("Severity", 12, "Choose a level from the dropdown. The levels themselves are editable on the 'Severity Scale' tab."),
    ("Platform", 16, "Original header: \"In which platform does the error occur?\""),
    ("Input Method", 20, "New field: device/method used when the issue occurred."),
    ("Screen Reader Used", 18, "New field: which screen reader (and version if relevant) was active, if any."),
    ("Status", 14, "New field: plain-text status - never rely on color alone."),
    ("Tester Name", 16, None),
    ("Notes", 30, "Original header: \"Provide any additional notes you think are relevant or helpful in regards to identifying or resolving the error.\" (merged with the old duplicate 'Notes' column)"),
]
SEVERITY_COL = 8   # column H
SECTION_COL = 3    # column C

SHEET_CONFIG = {
    "Visual": {"prefix": "V", "error_type_range": "VisualErrorTypes"},
    "Screen Reader": {"prefix": "SR", "error_type_range": "ScreenReaderErrorTypes"},
    "Keyboard": {"prefix": "KB", "error_type_range": "KeyboardErrorTypes"},
    "Controller": {"prefix": "C", "error_type_range": "ControllerErrorTypes"},
}

for sheet_name, cfg in SHEET_CONFIG.items():
    ws = reset_sheet(sheet_name)

    for idx, (label, width, comment_text) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=idx, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = HEADER_ALIGN
        c.protection = LOCKED
        if comment_text:
            c.comment = Comment(comment_text, COMMENT_AUTHOR)
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        # Issue ID only appears once a Game Section is chosen for this row -
        # keeps unused rows genuinely blank for screen-reader/table navigation.
        id_cell = ws.cell(row=r, column=1,
                          value=f'=IF(C{r}="","","{cfg["prefix"]}-"&TEXT(ROW()-1,"000"))')
        id_cell.font = BODY_FONT
        id_cell.alignment = Alignment(horizontal="left", vertical="top")
        id_cell.protection = LOCKED
        for idx in range(2, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=idx)
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
            cell.protection = UNLOCKED
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"

    add_table(ws, "tbl_" + cfg["prefix"],
              f"A1:{get_column_letter(len(COLUMNS))}{LAST_DATA_ROW}")

    def add_dv(formula, col_idx, sheet=ws):
        # showErrorMessage must be explicit: openpyxl defaults it to False, which
        # leaves the authored alert configured but never displayed, so free-typed
        # values are silently accepted and then quietly miss the Summary counts.
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            showErrorMessage=True)
        dv.error = "Please choose a value from the dropdown list."
        dv.errorTitle = "Invalid entry"
        letter = get_column_letter(col_idx)
        dv.add(f"{letter}{FIRST_DATA_ROW}:{letter}{LAST_DATA_ROW}")
        sheet.add_data_validation(dv)

    add_dv("GameSections", SECTION_COL)
    add_dv(cfg["error_type_range"], 4)
    add_dv("SeverityLevels", SEVERITY_COL)
    add_dv("PlatformList", 9)
    add_dv("InputMethodList", 10)
    add_dv("ScreenReaderList", 11)
    add_dv("StatusList", 12)

    protect_sheet(ws)

# ===========================================================================
# 5. Summary sheet
# ===========================================================================
summary = reset_sheet("Summary")
row_cursor = write_instructions(summary, [
    "Summary",
    "Totals below update automatically as you log issues on any tab - no need to refresh anything.",
    "Issues by Severity mirrors whatever levels are defined on the 'Severity Scale' tab. Issues by Status mirrors the Status list on the 'Setup' tab. Both count matching issues across every tab listed in Issues by Tab (bottom).",
    "How to add a custom tab (e.g. Adaptive Controllers, Mouse-Specific Controls):",
    "1. Right-click any existing log tab at the bottom of the workbook and choose 'Move or Copy'.",
    "2. Check 'Create a copy', choose a position for it (anywhere in the workbook - position doesn't matter), and click OK. This duplicates all its columns, dropdowns, and formatting.",
    "3. Right-click the new tab and choose 'Rename' - give it a clear name, e.g. 'Adaptive Controllers'.",
    "4. Come back here and type that exact tab name into the next blank row under 'Tab Name' below. Its totals start showing up automatically. A misspelled name shows 'Check tab name' instead of a count.",
    "Tip (optional): the new tab's Issue ID column reuses the prefix of the tab you copied (e.g. \"C-\" from Controller). To avoid duplicate-looking IDs, you can edit the formula in its first data cell to use your own prefix, then copy that cell down the column.",
])

# Layout is computed before writing so the severity totals can reference the
# per-tab working area further down the sheet.
sev_title_row = row_cursor
sev_header_row = sev_title_row + 1
sev_first_row = sev_header_row + 1
sev_last_row = sev_header_row + SEVERITY_ROW_COUNT

status_title_row = sev_last_row + 2
status_header_row = status_title_row + 1
status_first_row = status_header_row + 1
status_last_row = status_header_row + STATUS_ROW_COUNT

tab_title_row = status_last_row + 2
tab_header_row = tab_title_row + 1
TAB_NAME_SEED = ["Visual", "Screen Reader", "Keyboard", "Controller"]
TAB_ROW_COUNT = 20  # 4 seeded + growth room for custom tabs
tab_first_row = tab_header_row + 1
tab_last_row = tab_header_row + TAB_ROW_COUNT

# Working area: one column per severity/status slot, one row per tab. Each cell is
# an isolated single-sheet formula, so a mistyped tab name only breaks its own row
# instead of zeroing every total (which is what a single shared cross-sheet
# formula did). Severity/status totals are then a plain SUM down each column.
WORK_COL_START = 5  # column E
work_col = {i: get_column_letter(WORK_COL_START + i) for i in range(SEVERITY_ROW_COUNT)}
STATUS_WORK_COL_START = WORK_COL_START + SEVERITY_ROW_COUNT
status_work_col = {i: get_column_letter(STATUS_WORK_COL_START + i) for i in range(STATUS_ROW_COUNT)}

# --- Issues by Severity ---
summary.cell(row=sev_title_row, column=1, value="Issues by Severity").font = TITLE_FONT
summary.cell(row=sev_title_row, column=1).protection = LOCKED
write_header(summary, sev_header_row, ["Severity Value", "Meaning / Description", "Count"])

for i in range(SEVERITY_ROW_COUNT):
    r = sev_first_row + i
    scale_row = SEV_FIRST_DATA_ROW + i
    value_ref = f"'Severity Scale'!$A${scale_row}"
    meaning_ref = f"'Severity Scale'!$B${scale_row}"

    vc = summary.cell(row=r, column=1, value=f'=IF({value_ref}="","",{value_ref})')
    vc.font = BODY_FONT
    vc.protection = LOCKED

    mc = summary.cell(row=r, column=2, value=f'=IF({meaning_ref}="","",{meaning_ref})')
    mc.font = BODY_FONT
    mc.alignment = BODY_ALIGN
    mc.protection = LOCKED

    col = work_col[i]
    cc = summary.cell(
        row=r, column=3,
        value=f'=IF({value_ref}="","",SUM({col}{tab_first_row}:{col}{tab_last_row}))'
    )
    cc.font = BODY_FONT
    cc.protection = LOCKED

add_table(summary, "tbl_SummarySeverity", f"A{sev_header_row}:C{sev_last_row}")

# --- Issues by Status ---
summary.cell(row=status_title_row, column=1, value="Issues by Status").font = TITLE_FONT
summary.cell(row=status_title_row, column=1).protection = LOCKED
write_header(summary, status_header_row, ["Status", "Count"])

for i in range(STATUS_ROW_COUNT):
    r = status_first_row + i
    status_ref = f"'Setup'!$A${STATUS_FIRST_DATA_ROW + i}"

    vc = summary.cell(row=r, column=1, value=f'=IF({status_ref}="","",{status_ref})')
    vc.font = BODY_FONT
    vc.protection = LOCKED

    col = status_work_col[i]
    cc = summary.cell(
        row=r, column=2,
        value=f'=IF({status_ref}="","",SUM({col}{tab_first_row}:{col}{tab_last_row}))'
    )
    cc.font = BODY_FONT
    cc.protection = LOCKED

add_table(summary, "tbl_SummaryStatus", f"A{status_header_row}:B{status_last_row}")

# --- Issues by Tab ---
summary.cell(row=tab_title_row, column=1, value="Issues by Tab").font = TITLE_FONT
summary.cell(row=tab_title_row, column=1).protection = LOCKED
write_header(summary, tab_header_row, ["Tab Name", "Total Issues"])

work_label = summary.cell(row=tab_title_row, column=WORK_COL_START,
                          value="Working area - calculations, do not edit")
work_label.font = TITLE_FONT
work_label.protection = LOCKED
write_header(summary, tab_header_row,
             [f"Level {i + 1}" for i in range(SEVERITY_ROW_COUNT)],
             start_col=WORK_COL_START)
write_header(summary, tab_header_row,
             [f"Status {i + 1}" for i in range(STATUS_ROW_COUNT)],
             start_col=STATUS_WORK_COL_START)

for i in range(TAB_ROW_COUNT):
    r = tab_first_row + i
    nc = summary.cell(row=r, column=1,
                      value=TAB_NAME_SEED[i] if i < len(TAB_NAME_SEED) else None)
    nc.font = BODY_FONT
    nc.protection = UNLOCKED

    # ISREF validates the typed tab name first: COUNTA/COUNTIFS over an INDIRECT to a
    # nonexistent sheet do NOT raise a catchable error (they treat the #REF! itself as
    # one item), so IFERROR alone cannot detect a mistyped tab name - ISREF can.
    # Apostrophes are legal in sheet names ("Ava's Tab") and must be doubled inside a
    # quoted sheet reference, or INDIRECT resolves to #REF! and the row wrongly reads
    # as a mistyped tab name.
    tab_ref = f'SUBSTITUTE($A{r},"\'","\'\'")'
    is_valid = f'ISREF(INDIRECT("\'"&{tab_ref}&"\'!$C$2"))'
    section_range = f'INDIRECT("\'"&{tab_ref}&"\'!$C$2:$C${LAST_DATA_ROW}")'
    severity_range = f'INDIRECT("\'"&{tab_ref}&"\'!$H$2:$H${LAST_DATA_ROW}")'
    status_range = f'INDIRECT("\'"&{tab_ref}&"\'!$L$2:$L${LAST_DATA_ROW}")'

    # Counts column C (Game Section), not column A (Issue ID): column A holds a formula
    # on every row, and COUNTA treats a formula returning "" as non-blank, which would
    # count every unused row too.
    total_cell = summary.cell(
        row=r, column=2,
        value=(f'=IF($A{r}="","",IF({is_valid},COUNTA({section_range}),"Check tab name"))')
    )
    total_cell.font = BODY_FONT
    total_cell.protection = LOCKED

    for j in range(SEVERITY_ROW_COUNT):
        value_ref = f"'Severity Scale'!$A${SEV_FIRST_DATA_ROW + j}"
        # Blank (not text) when unusable, so the SUM above ignores it cleanly.
        wc = summary.cell(
            row=r, column=WORK_COL_START + j,
            value=(
                f'=IF(OR($A{r}="",{value_ref}=""),"",'
                f'IF({is_valid},COUNTIFS({section_range},"<>",{severity_range},{value_ref}),""))'
            )
        )
        wc.font = BODY_FONT
        wc.protection = LOCKED

    for k in range(STATUS_ROW_COUNT):
        status_value_ref = f"'Setup'!$A${STATUS_FIRST_DATA_ROW + k}"
        swc = summary.cell(
            row=r, column=STATUS_WORK_COL_START + k,
            value=(
                f'=IF(OR($A{r}="",{status_value_ref}=""),"",'
                f'IF({is_valid},COUNTIFS({section_range},"<>",{status_range},{status_value_ref}),""))'
            )
        )
        swc.font = BODY_FONT
        swc.protection = LOCKED

add_table(summary, "tbl_SummaryTabs", f"A{tab_header_row}:B{tab_last_row}")

summary.column_dimensions["A"].width = 26
summary.column_dimensions["B"].width = 14
summary.column_dimensions["C"].width = 12
for col in work_col.values():
    summary.column_dimensions[col].width = 9
for col in status_work_col.values():
    summary.column_dimensions[col].width = 9
summary.freeze_panes = "A2"
protect_sheet(summary)

# ===========================================================================
# 6. Sheet order
# ===========================================================================
order = ["Setup", "Game Sections", "Severity Scale",
         "Visual", "Screen Reader", "Keyboard", "Controller", "Summary"]
wb._sheets = [wb[name] for name in order]
wb.active = 0

# Authorship is deliberate: this template is published openly and credited to its
# creator. Contact details are intentionally NOT included - a name is attribution,
# an email address in a public file is a spam magnet.
wb.properties.creator = AUTHOR_NAME
wb.properties.lastModifiedBy = AUTHOR_NAME
wb.properties.title = "Game Accessibility Testing Template"
wb.properties.description = None
wb.properties.keywords = None
wb.properties.category = None
wb.properties.identifier = None
wb.properties.language = None
wb.properties.subject = None

wb.save(SRC)
print("Saved:", SRC)
